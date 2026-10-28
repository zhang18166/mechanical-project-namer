# SPDX-License-Identifier: GPL-3.0-only
# Copyright (C) 2026 zhang18166

import csv
import json
import logging
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import QUrl
from PyQt6.QtPdf import QPdfDocument

from mechanical_project_namer.ai import AITaskState, AITaskType, NamingMemory, StreamEvent
from mechanical_project_namer.data import DataExporter, Database, Repository
from mechanical_project_namer.logging_setup import solidworks_logger
from mechanical_project_namer.projects import ProjectService
from mechanical_project_namer.solidworks.models import OperationRequest, OperationResult
from mechanical_project_namer.solidworks.worker_entry import run_worker
from mechanical_project_namer.ui.common import open_local_path
from mechanical_project_namer.ui.design_chat_dialog import DesignChatDialog
from mechanical_project_namer.ui.shell import ApplicationShell


@pytest.mark.parametrize("factory", [str, Path])
def test_local_path_service_uses_qurl_for_string_and_path(tmp_path, factory):
    target = tmp_path / "中文目录"
    target.mkdir()
    received = []
    result = open_local_path(factory(target), "directory", opener=lambda url: received.append(url) or True)
    assert result.ok and len(received) == 1
    assert isinstance(received[0], QUrl)
    assert Path(received[0].toLocalFile()) == target


@pytest.mark.parametrize(
    ("value", "expected_type", "code"),
    [("", "any", "empty"), (None, "file", "empty"), (123, "any", "invalid_type")],
)
def test_local_path_service_rejects_empty_and_wrong_value(value, expected_type, code):
    assert open_local_path(value, expected_type, opener=lambda _url: True).code == code


def test_local_path_service_handles_missing_wrong_type_and_open_failure(tmp_path):
    missing = open_local_path(tmp_path / "不存在.log", "file", opener=lambda _url: True)
    wrong = open_local_path(tmp_path, "file", opener=lambda _url: True)
    failed = open_local_path(tmp_path, "directory", opener=lambda _url: False)
    assert (missing.code, wrong.code, failed.code) == ("missing", "wrong_type", "open_failed")


def test_local_path_service_expands_environment_and_chinese_file(tmp_path, monkeypatch):
    target = tmp_path / "压弧滑座 日志.txt"
    target.write_text("中文", encoding="utf-8")
    monkeypatch.setenv("V141_PATH_ROOT", str(tmp_path))
    received = []
    result = open_local_path(r"%V141_PATH_ROOT%\压弧滑座 日志.txt", "file", opener=lambda url: received.append(url) or True)
    assert result.ok and Path(received[0].toLocalFile()) == target


class _FakeBridge:
    def execute(self, request, progress):
        progress("connected", "已连接SOLIDWORKS 2025 SP5")
        progress("template_resolved", r"模板：D:\中文模板\零件模板.prtdot")
        progress("saving", r"正在保存：D:\2026年\8月份\ZGQ-03-005 压弧滑座.SLDPRT")
        return OperationResult(request.operation_id, True, "file_saved", "文件验证成功")


def test_solidworks_worker_events_are_structured_utf8(tmp_path, monkeypatch):
    result_path = tmp_path / "result.json"
    request = OperationRequest("utf8-operation", action="probe", result_path=str(result_path))
    request_path = tmp_path / "request.json"
    request_path.write_text(json.dumps(request.to_dict(), ensure_ascii=False), encoding="utf-8")
    events = []
    monkeypatch.setattr("mechanical_project_namer.solidworks.worker_entry._emit_stdout", events.append)
    assert run_worker(str(request_path), _FakeBridge()) == 0
    progress = [item for item in events if item.get("type") == "progress"]
    assert all(item["operation_id"] == "utf8-operation" for item in progress)
    assert all("T" in item["timestamp"] and item["data"] == {} for item in progress)
    text = json.dumps(progress, ensure_ascii=False)
    assert "已连接SOLIDWORKS" in text and "中文模板" in text and "压弧滑座.SLDPRT" in text


def test_solidworks_result_accepts_forward_compatible_error_details():
    result = OperationResult.from_dict({
        "operation_id": "failure", "success": False, "stage": "save",
        "message": "保存失败", "error_details": {"path": "中文路径"},
    })
    assert not result.success and result.message == "保存失败"


def test_dedicated_solidworks_log_is_utf8_and_rolling(tmp_path, monkeypatch):
    logger = logging.getLogger("mechanical_project_namer.solidworks.operations")
    for handler in tuple(logger.handlers):
        handler.close(); logger.removeHandler(handler)
    monkeypatch.setattr("mechanical_project_namer.logging_setup.logs_dir", lambda: tmp_path)
    sw_logger = solidworks_logger()
    sw_logger.info("中文模板：D:/中文模板/零件模板.prtdot")
    for handler in sw_logger.handlers:
        handler.flush()
    target = tmp_path / "solidworks.log"
    assert target.is_file() and "中文模板" in target.read_text(encoding="utf-8")
    assert sw_logger.handlers[0].backupCount == 5


@pytest.fixture
def v141_repository(tmp_path):
    database = Database(tmp_path / "v141.db")
    database.migrate()
    repository = Repository(database)
    project = ProjectService(repository).create("圆管自动压圈打圈机", "送料主动轴", "ZGQ", directory=str(tmp_path))
    return database, repository, project


def test_conversation_second_request_contains_first_round_and_restart(v141_repository):
    database, repository, project = v141_repository
    conversation_id = repository.create_conversation(project.id, "新对话", "deepseek-v4-flash", False, "high")
    repository.add_ai_message(conversation_id, "user", "我正在45钢和40Cr之间选择，用于一个送料主动轴。")
    repository.add_ai_message(conversation_id, "assistant", "已了解这两个候选材料。")
    reopened = Repository(Database(database.path))
    context = reopened.prepare_conversation_context(conversation_id)
    assert [item["role"] for item in context["recent_messages"]] == ["user", "assistant"]
    assert "45钢和40Cr" in context["recent_messages"][0]["content"]
    assert reopened.get_conversation(conversation_id)["title"].startswith("我正在45钢和40Cr")


def test_new_conversation_does_not_inherit_previous_history(v141_repository):
    _, repository, project = v141_repository
    old_id = repository.create_conversation(project.id, "材料", "m", False, "high")
    repository.add_ai_message(old_id, "user", "45钢和40Cr")
    new_id = repository.create_conversation(project.id, "新对话", "m", False, "high")
    assert repository.prepare_conversation_context(new_id)["recent_messages"] == []


def test_conversation_summary_bounds_recent_history(v141_repository):
    _, repository, project = v141_repository
    conversation_id = repository.create_conversation(project.id, "长会话", "m", False, "high")
    for index in range(24):
        repository.add_ai_message(conversation_id, "user" if index % 2 == 0 else "assistant", f"第{index}条消息")
    context = repository.prepare_conversation_context(conversation_id, recent_limit=8, summary_threshold=12)
    assert context["summary_loaded"] and "第0条消息" in context["summary"]
    assert len(context["recent_messages"]) == 8
    assert repository.get_conversation(conversation_id)["summary_through_message_id"] > 0


def test_design_chat_message_builder_keeps_history_with_all_memory_switches_off(v141_repository):
    database, repository, project = v141_repository
    conversation_id = repository.create_conversation(project.id, "材料", "m", False, "high")
    repository.add_ai_message(conversation_id, "user", "我在45钢和40Cr之间选择，用于送料主动轴。")
    repository.add_ai_message(conversation_id, "assistant", "请继续。")
    shell = object.__new__(ApplicationShell)
    shell.repository = repository
    shell.memory = NamingMemory(database)
    shell.chat_conversation_id = conversation_id
    shell.current_project_id = project.id
    shell.current_module_code = "02"
    messages, stats = shell._build_design_chat_messages("我刚才说的是哪两种材料？", {
        "use_project": False, "use_history": False, "use_naming": False, "use_memory": False,
    })
    assert [item["role"] for item in messages] == ["system", "user", "assistant", "user"]
    assert "45钢和40Cr" in messages[1]["content"] and stats["history_messages"] == 2


def test_conversation_list_refresh_select_and_delete(qtbot, v141_repository):
    _, repository, project = v141_repository
    first = repository.create_conversation(project.id, "第一会话", "m", False, "high")
    second = repository.create_conversation(project.id, "第二会话", "m", False, "high")
    dialog = DesignChatDialog(repository.list_conversations(project.id)); qtbot.addWidget(dialog)
    dialog.set_conversations(repository.list_conversations(project.id), first)
    assert dialog.selected_conversation_id() == first
    repository.rename_conversation(first, "送料轴讨论")
    repository.delete_conversation(second)
    dialog.set_conversations(repository.list_conversations(project.id), first)
    assert dialog.conversation_list.count() == 1 and dialog.conversation_list.item(0).text() == "送料轴讨论"


def test_ai_task_types_have_independent_cancel_tokens():
    class Worker:
        def __init__(self): self.cancelled = False; self.cancel_event = object()
        def cancel(self): self.cancelled = True
    planning_worker, chat_worker = Worker(), Worker()
    planning = AITaskState(AITaskType.AI_PLANNING, "planning", planning_worker, object())
    chat = AITaskState(AITaskType.AI_CHAT, "chat", chat_worker, object())
    planning.cancel()
    assert planning_worker.cancelled and not chat_worker.cancelled
    assert planning.task_id != chat.task_id and planning.cancel_token is not chat.cancel_token


@pytest.fixture
def export_project(v141_repository):
    _, repository, project = v141_repository
    for index in range(1, 61):
        repository.save_named_record(project.id, "03", "自制机加工件", f"压弧滑座{index}", f"备注,{index}\n第二行")
    repository.set_record_status(repository.list_records(project.id)[0].id, "void")
    return repository, project


def test_pdf_excel_csv_exports_are_consistent_atomic_and_chinese(qtbot, export_project, tmp_path):
    repository, project = export_project
    exporter = DataExporter(repository, tmp_path)
    csv_result = exporter.export_records("csv", tmp_path / "记录.csv", project_id=project.id)
    xlsx_result = exporter.export_records("xlsx", tmp_path / "记录.xlsx", project_id=project.id)
    pdf_result = exporter.export_records("pdf", tmp_path / "记录.pdf", project_id=project.id)
    assert csv_result.record_count == xlsx_result.record_count == pdf_result.record_count == 60
    assert csv_result.path.read_bytes().startswith(b"\xef\xbb\xbf")
    with csv_result.path.open("r", encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))
    assert csv_rows[0] == list(__import__("mechanical_project_namer.data.exporter", fromlist=["RECORD_HEADERS"]).RECORD_HEADERS)
    assert "压弧滑座" in "".join(csv_rows[1])
    workbook = load_workbook(xlsx_result.path)
    assert workbook.sheetnames == ["项目概览", "项目文件记录", "编号状态"]
    assert workbook["项目文件记录"].freeze_panes == "A2"
    assert workbook["项目文件记录"].auto_filter.ref
    assert workbook["项目文件记录"].max_row == 61
    workbook.close()
    pdf = QPdfDocument(None); pdf.load(str(pdf_result.path))
    assert pdf.pageCount() == pdf_result.page_count and pdf.pageCount() > 1
    pdf.close()
    assert not list(tmp_path.glob(".*.tmp.*"))


def test_export_scope_empty_project_selected_status_and_no_silent_overwrite(qtbot, v141_repository, tmp_path):
    _, repository, project = v141_repository
    exporter = DataExporter(repository, tmp_path)
    empty = exporter.export_records("xlsx", tmp_path / "empty.xlsx", project_id=project.id)
    assert empty.record_count == 0
    first = repository.save_named_record(project.id, "03", "自制机加工件", "滑座")
    second = repository.save_named_record(project.id, "03", "自制机加工件", "轴座")
    repository.set_record_status(second.id, "void")
    selected = exporter.export_records("csv", tmp_path / "selected.csv", project_id=project.id, record_ids={first.id})
    void = exporter.export_records("csv", tmp_path / "void.csv", project_id=project.id, status="void")
    assert selected.record_count == void.record_count == 1
    with pytest.raises(FileExistsError):
        exporter.export_records("csv", selected.path, project_id=project.id)


def test_exports_do_not_contain_api_credentials(qtbot, export_project, tmp_path):
    repository, project = export_project
    exporter = DataExporter(repository, tmp_path)
    paths = [exporter.export_records(fmt, tmp_path / f"safe.{fmt}", project_id=project.id).path for fmt in ("csv", "xlsx", "pdf")]
    for path in paths:
        raw = path.read_bytes()
        assert b"Authorization" not in raw and b"sk-" not in raw
