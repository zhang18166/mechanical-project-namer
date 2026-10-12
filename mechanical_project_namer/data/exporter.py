# SPDX-License-Identifier: GPL-3.0-only
# Copyright (C) 2026 zhang18166

import csv
import json
import os
import tempfile
import textwrap
import unicodedata
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path

from PyQt6.QtCore import QMarginsF, QRect, Qt
from PyQt6.QtGui import QColor, QFont, QFontMetrics, QPageLayout, QPageSize, QPainter, QPdfWriter

from ..paths import exports_dir
from ..safe_fs import remove_file
from .. import __version__
from .repository import Repository


STATUS_NAMES = {"used": "已使用", "void": "已作废", "reserved": "已预留"}

RECORD_HEADERS = (
    "序号", "日期", "完整名称", "项目代号", "模块编号", "模块名称", "零件序号", "零件名称",
    "零件类别", "额外备注", "状态", "物理文件路径", "文件是否存在", "创建方式", "创建时间", "更新时间",
    "文件大小", "创建来源",
)


@dataclass(frozen=True)
class RecordExportResult:
    format: str
    path: Path
    record_count: int
    project_name: str
    project_code: str
    sheets: tuple[str, ...] = ()
    page_count: int = 0
    encoding: str = ""


class DataExporter:
    def __init__(self, repository: Repository, directory: Path | None = None) -> None:
        self.repository = repository
        self.directory = Path(directory or exports_dir())
        self._snapshot_token = None
        self._state_token = None
        self._schema_written = False

    def _database_token(self) -> tuple:
        # Include WAL: SQLite commits need not update the main database file.
        database = self.repository.database.path
        token = []
        for path in (database, Path(str(database) + "-wal")):
            try:
                stat = path.stat()
                token.append((stat.st_mtime_ns, stat.st_size))
            except FileNotFoundError:
                token.append(None)
        return tuple(token)

    def export_all(self, current_project_id: int | None = None) -> None:
        self.directory.mkdir(parents=True, exist_ok=True)
        token = self._database_token()
        snapshot_changed = token != self._snapshot_token or not (self.directory / "项目数据快照.json").is_file()
        state_changed = (token, current_project_id) != self._state_token or not (self.directory / "当前项目状态.json").is_file()
        schema_changed = not self._schema_written or not (self.directory / "数据格式说明.json").is_file()
        if not (snapshot_changed or state_changed or schema_changed):
            return
        updated = datetime.now().astimezone().isoformat(timespec="seconds")
        if snapshot_changed:
            projects = self.repository.list_projects(include_archived=True)
            records = self.repository.list_records()
            snapshot = {
                "schema_version": __version__, "updated_at": updated, "read_only": True,
                "projects": [asdict(project) for project in projects],
                "records": [asdict(record) for record in records],
            }
            self._atomic_json("项目数据快照.json", snapshot)
            self._snapshot_token = token
        if state_changed:
            self._atomic_json("当前项目状态.json", self._current_state(current_project_id, updated))
            self._state_token = (token, current_project_id)
        if schema_changed:
            self._atomic_json("数据格式说明.json", self._schema_description(updated))
            self._schema_written = True

    def export_csv(self, target: Path, project_id: int | None = None) -> int:
        result = self.export_records("csv", target, project_id=project_id, overwrite=True)
        return result.record_count

    def export_records(
        self,
        file_format: str,
        target: Path,
        *,
        project_id: int | None,
        query: str = "",
        date_from: str = "",
        date_to: str = "",
        record_ids: set[int] | None = None,
        status: str = "",
        overwrite: bool = False,
    ) -> RecordExportResult:
        file_format = file_format.lower().lstrip(".")
        if file_format not in {"csv", "xlsx", "pdf"}:
            raise ValueError("不支持的导出格式。")
        target = Path(target)
        expected_suffix = ".xlsx" if file_format == "xlsx" else f".{file_format}"
        if target.suffix.lower() != expected_suffix:
            target = target.with_suffix(expected_suffix)
        if target.exists() and not overwrite:
            raise FileExistsError(f"目标文件已存在：{target}")
        target.parent.mkdir(parents=True, exist_ok=True)
        records = self.repository.list_records(project_id, query, date_from, date_to)
        if record_ids is not None:
            records = [record for record in records if record.id in record_ids]
        if status:
            records = [record for record in records if record.status == status]
        modules = {}
        if project_id is not None:
            modules = {item["module_code"]: item["module_name"] for item in self.repository.list_modules(project_id, enabled_only=False)}
        rows = self._record_rows(records, modules)
        project = self.repository.get_project(project_id) if project_id is not None else None
        fd, temporary_name = tempfile.mkstemp(
            prefix=f".{target.stem}.tmp.", suffix=target.suffix, dir=target.parent
        )
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            if file_format == "csv":
                self._write_csv(temporary, rows)
                self._validate_csv(temporary, len(rows))
                metadata = {"encoding": "UTF-8-SIG"}
            elif file_format == "xlsx":
                self._write_excel(temporary, project, records, rows, modules)
                sheets = self._validate_excel(temporary, len(rows))
                metadata = {"sheets": tuple(sheets)}
            else:
                pages = self._write_pdf(temporary, project, rows)
                self._validate_pdf(temporary, pages)
                metadata = {"page_count": pages}
            os.replace(temporary, target)
        finally:
            remove_file(temporary, what="导出临时文件")
        return RecordExportResult(
            file_format.upper() if file_format != "xlsx" else "Excel",
            target,
            len(rows),
            project.name if project else "全部项目",
            project.code if project else "ALL",
            **metadata,
        )

    @staticmethod
    def _record_rows(records, modules: dict[str, str]) -> list[list]:
        rows = []
        for index, record in enumerate(records, 1):
            rows.append([
                index, record.record_date, record.display_name, record.project_code,
                record.module_code, modules.get(record.module_code, ""), f"{record.part_sequence:03d}",
                record.part_name, record.part_type, record.extra_notes,
                STATUS_NAMES.get(record.status, record.status), record.file_path,
                "是" if record.file_exists and record.file_path and Path(record.file_path).is_file() else "否",
                record.creation_method or record.creation_source, record.created_at, record.updated_at,
                record.file_size, record.creation_source,
            ])
        return rows

    @staticmethod
    def _write_csv(target: Path, rows: list[list]) -> None:
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(RECORD_HEADERS)
            writer.writerows(rows)
            stream.flush()
            os.fsync(stream.fileno())

    @staticmethod
    def _validate_csv(target: Path, expected_rows: int) -> None:
        if not target.read_bytes().startswith(b"\xef\xbb\xbf"):
            raise ValueError("CSV不是UTF-8-SIG编码。")
        with target.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        if not rows or tuple(rows[0]) != RECORD_HEADERS or len(rows) - 1 != expected_rows:
            raise ValueError("CSV回读验证失败。")

    def _write_excel(self, target: Path, project, records, rows: list[list], modules: dict[str, str]) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        overview = workbook.active
        overview.title = "项目概览"
        now = datetime.now().astimezone()
        statuses = {"normal": "正常", "archived": "已归档", "trash": "回收站"}
        project_types = {"formal": "正式项目", "test": "测试项目"}
        used = sum(record.status == "used" for record in records)
        void = sum(record.status == "void" for record in records)
        overview_rows = (
            ("项目名称", project.name if project else "全部项目"),
            ("项目代号", project.code if project else "ALL"),
            ("项目状态", statuses.get(project.project_status, project.project_status) if project else "全部"),
            ("项目类型", project_types.get(project.project_type, project.project_type) if project else "全部"),
            ("项目目录", project.directory if project else ""),
            ("模块数量", len(modules)), ("零件记录数量", len(records)),
            ("已使用数量", used), ("已作废数量", void),
            ("创建时间", self._excel_datetime(project.created_at) if project else None),
            ("最后使用时间", self._excel_datetime(project.last_used_at) if project else None),
            ("导出时间", now.replace(tzinfo=None)), ("软件版本", f"V{__version__}"),
        )
        for row in overview_rows:
            overview.append(row)
        overview.column_dimensions["A"].width = 20
        overview.column_dimensions["B"].width = 70
        for row in range(1, overview.max_row + 1):
            overview.cell(row, 1).font = Font(bold=True)
            overview.cell(row, 1).fill = PatternFill("solid", fgColor="EAF2F8")
        for row in range(10, 13):
            overview.cell(row, 2).number_format = "yyyy-mm-dd hh:mm:ss"

        records_sheet = workbook.create_sheet("项目文件记录")
        records_sheet.append(RECORD_HEADERS)
        for row in rows:
            converted = list(row)
            converted[1] = self._excel_date(converted[1])
            converted[14] = self._excel_datetime(converted[14])
            converted[15] = self._excel_datetime(converted[15])
            records_sheet.append(converted)
        self._format_tabular_sheet(records_sheet, wrap_columns={3, 8, 10, 12})
        for row in range(2, records_sheet.max_row + 1):
            records_sheet.cell(row, 2).number_format = "yyyy-mm-dd"
            records_sheet.cell(row, 5).number_format = "@"
            records_sheet.cell(row, 7).number_format = "@"
            records_sheet.cell(row, 15).number_format = "yyyy-mm-dd hh:mm:ss"
            records_sheet.cell(row, 16).number_format = "yyyy-mm-dd hh:mm:ss"

        number_sheet = workbook.create_sheet("编号状态")
        number_headers = ("模块编号", "模块名称", "零件类别", "编号范围", "已使用数量", "下一个可用编号")
        number_sheet.append(number_headers)
        ranges = {"总成": "000", "自制机加工件": "001-099", "钣金/焊接件": "101-199", "采购/标准件": "201-299"}
        counts = {}
        for record in records:
            counts[(record.module_code, record.part_type)] = counts.get((record.module_code, record.part_type), 0) + 1
        if project:
            for module in self.repository.project_numbering_status(project.id):
                for item in module["types"]:
                    number_sheet.append([
                        module["module_code"], module["module_name"], item["part_type"],
                        ranges.get(item["part_type"], ""), counts.get((module["module_code"], item["part_type"]), 0),
                        item["next_number"],
                    ])
        self._format_tabular_sheet(number_sheet)
        for row in range(2, number_sheet.max_row + 1):
            number_sheet.cell(row, 1).number_format = "@"
            number_sheet.cell(row, 6).number_format = "@"
        workbook.save(target)

    @staticmethod
    def _format_tabular_sheet(sheet, wrap_columns: set[int] | None = None) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wrap_columns = wrap_columns or set()
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_columns)
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            def display_width(value: str) -> int:
                return sum(2 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 1 for character in value)
            width = min(60, max(10, max((display_width(value) for value in values), default=10) + 2))
            sheet.column_dimensions[get_column_letter(column)].width = width

    @staticmethod
    def _validate_excel(target: Path, expected_rows: int) -> list[str]:
        from openpyxl import load_workbook

        workbook = load_workbook(target, read_only=True, data_only=True)
        try:
            required = ["项目概览", "项目文件记录", "编号状态"]
            if any(name not in workbook.sheetnames for name in required):
                raise ValueError("Excel缺少必需工作表。")
            if workbook["项目文件记录"].max_row - 1 != expected_rows:
                raise ValueError("Excel记录数量验证失败。")
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    @staticmethod
    def _excel_date(value):
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return value

    @staticmethod
    def _excel_datetime(value):
        if not value:
            return None
        try:
            return datetime.fromisoformat(str(value)).replace(tzinfo=None)
        except ValueError:
            return value

    @staticmethod
    def _write_pdf(target: Path, project, rows: list[list]) -> int:
        writer = QPdfWriter(str(target))
        writer.setResolution(144)
        writer.setPageLayout(QPageLayout(
            QPageSize(QPageSize.PageSizeId.A4), QPageLayout.Orientation.Landscape,
            QMarginsF(10, 10, 10, 10), QPageLayout.Unit.Millimeter,
        ))
        painter = QPainter(writer)
        if not painter.isActive():
            raise OSError("无法创建PDF绘图设备。")
        page = writer.pageLayout().paintRectPixels(writer.resolution())
        title_font = QFont("Microsoft YaHei UI", 13, QFont.Weight.Bold)
        text_font = QFont("Microsoft YaHei UI", 7)
        header_font = QFont("Microsoft YaHei UI", 7, QFont.Weight.Bold)
        margin, title_height, footer_height, header_height = 8, 70, 30, 42
        table_headers = ("日期", "完整名称", "模块", "零件名称", "零件类别", "额外备注", "状态", "文件路径")
        indexes = (1, 2, 4, 7, 8, 9, 10, 11)
        weights = (8, 20, 6, 12, 9, 14, 7, 24)
        available_width = page.width() - margin * 2
        widths = [int(available_width * weight / sum(weights)) for weight in weights]
        widths[-1] += available_width - sum(widths)
        metrics = QFontMetrics(text_font, writer)

        def row_height(values):
            heights = []
            for value, width in zip(values, widths):
                bounds = metrics.boundingRect(
                    QRect(0, 0, max(1, width - 6), 1000),
                    int(Qt.TextFlag.TextWordWrap),
                    str(value or ""),
                )
                explicit_lines = str(value or "").count("\n") + 1
                heights.append(max(bounds.height() + 8, explicit_lines * metrics.lineSpacing() + 8))
            return max(28, min(110, max(heights, default=28)))

        pdf_rows = []
        for row in rows:
            values = [row[index] for index in indexes]
            values[-1] = "\n".join(textwrap.wrap(str(values[-1] or ""), width=38, break_long_words=True, break_on_hyphens=False))
            pdf_rows.append((values, row_height(values)))
        pages: list[list] = []
        current, used_height = [], title_height + header_height
        body_limit = page.height() - footer_height
        for values, height in pdf_rows:
            if current and used_height + height > body_limit:
                pages.append(current); current = []; used_height = title_height + header_height
            current.append((values, height)); used_height += height
        pages.append(current)
        total_pages = max(1, len(pages))
        for page_index, page_rows in enumerate(pages):
            if page_index:
                writer.newPage()
            painter.setFont(title_font)
            title = f"{project.name} · {project.code} 项目文件记录" if project else "全部项目文件记录"
            painter.drawText(QRect(margin, margin, available_width, 32), int(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter), title)
            painter.setFont(text_font)
            exported = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S")
            painter.drawText(QRect(margin, margin + 32, available_width, 26), int(Qt.AlignmentFlag.AlignLeft), f"导出时间：{exported}    记录数量：{len(rows)}")
            y = margin + title_height
            painter.setFont(header_font); painter.setBrush(QColor("#D9EAF7")); painter.setPen(QColor("#536471"))
            x = margin
            for header, width in zip(table_headers, widths):
                rect = QRect(x, y, width, header_height); painter.fillRect(rect, QColor("#D9EAF7")); painter.drawRect(rect)
                painter.drawText(rect.adjusted(3, 2, -3, -2), int(Qt.AlignmentFlag.AlignCenter | Qt.TextFlag.TextWordWrap), header); x += width
            y += header_height; painter.setFont(text_font); painter.setBrush(Qt.BrushStyle.NoBrush)
            for row_index, (values, height) in enumerate(page_rows):
                x = margin
                if row_index % 2:
                    painter.fillRect(QRect(x, y, available_width, height), QColor("#F6F8FA"))
                for value, width in zip(values, widths):
                    rect = QRect(x, y, width, height); painter.drawRect(rect)
                    painter.drawText(
                        rect.adjusted(3, 3, -3, -3),
                        int(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignTop | Qt.TextFlag.TextWordWrap),
                        str(value or ""),
                    ); x += width
                y += height
            painter.drawText(QRect(margin, page.height() - footer_height, available_width, 22), int(Qt.AlignmentFlag.AlignRight), f"第 {page_index + 1} / {total_pages} 页")
        painter.end()
        return total_pages

    @staticmethod
    def _validate_pdf(target: Path, expected_pages: int) -> None:
        from PyQt6.QtPdf import QPdfDocument

        raw = target.read_bytes()
        if not raw.startswith(b"%PDF-") or len(raw) == 0:
            raise ValueError("PDF文件头验证失败。")
        document = QPdfDocument(None)
        document.load(str(target))
        try:
            if document.pageCount() <= 0 or document.pageCount() != expected_pages:
                raise ValueError("PDF页数验证失败。")
        finally:
            document.close()

    def export_project_package(self, project_id: int, target: Path) -> Path:
        project = self.repository.get_project(project_id)
        properties = self.repository.project_properties(project_id)
        target = Path(target)
        if target.suffix.lower() != ".zip":
            target = target.with_suffix(".zip")
        target.parent.mkdir(parents=True, exist_ok=True)
        modules = self.repository.list_modules(project_id, enabled_only=False)
        records = self.repository.list_records(project_id)
        plans = self.repository.list_ai_plans(project_id, limit=10000)
        memories = self.repository.list_project_memories(project_id, limit=10000)
        with self.repository.database.connect() as connection:
            feedback = [dict(row) for row in connection.execute(
                "SELECT id,module_code,ai_name,final_name,action,reject_reason,input_description,confidence,use_count,created_at,updated_at FROM ai_feedback WHERE project_id=? ORDER BY updated_at",
                (project_id,),
            ).fetchall()]
        project_data = asdict(project)
        project_data["statistics"] = {key: value for key, value in properties.items() if key != "project"}
        readme = (
            f"# {project.name} ({project.code})\n\n"
            "本资料包仅包含机械项目命名助手中的项目管理数据。\n"
            "不包含 SolidWorks、STEP、NX、PDF、NC 等真实设计文件，也不包含 API Key 或 Windows 凭据。\n"
        )
        fd, temporary = tempfile.mkstemp(prefix=f".{target.stem}.", suffix=".zip.tmp", dir=target.parent)
        os.close(fd)
        try:
            with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as bundle:
                for filename, payload in (
                    ("project.json", project_data), ("modules.json", modules),
                    ("parts.json", [asdict(record) for record in records]),
                    ("ai_plans.json", plans), ("project_memories.json", memories),
                ):
                    bundle.writestr(filename, json.dumps(payload, ensure_ascii=False, indent=2, default=str))
                examples = "\n".join(json.dumps(item, ensure_ascii=False, default=str) for item in feedback)
                bundle.writestr("naming_examples.jsonl", examples + ("\n" if examples else ""))
                bundle.writestr("README.md", readme)
                csv_lines = [["完整名称", "零件名称", "模块", "类别", "状态", "创建时间", "物理文件路径", "文件大小", "创建来源"]]
                csv_lines.extend([[r.display_name, r.part_name, r.module_code, r.part_type, STATUS_NAMES.get(r.status, r.status), r.created_at, r.file_path, r.file_size, r.creation_source] for r in records])
                import io
                stream = io.StringIO(newline="")
                csv.writer(stream).writerows(csv_lines)
                bundle.writestr("records.csv", "\ufeff" + stream.getvalue())
            os.replace(temporary, target)
        finally:
            remove_file(temporary, what="导出临时文件")
        self.repository.audit_project_action(project_id, "export")
        return target

    def _current_state(self, project_id: int | None, updated: str) -> dict:
        if project_id is None:
            return {"schema_version": __version__, "current_project_id": None, "updated_at": updated}
        try:
            project = self.repository.get_project(project_id)
        except LookupError:
            return {"schema_version": __version__, "current_project_id": None, "updated_at": updated}
        return {
            "schema_version": __version__,
            "current_project_id": project.id,
            "current_project_name": project.name,
            "current_project_code": project.code,
            "numbering_rule": "项目代号-模块编号-零件序号；各模块和零件类别独立计数",
            "modules": self.repository.list_modules(project.id, enabled_only=False),
            "latest_created_number": project.latest_number,
            "project_directory": project.directory,
            "solidworks_part_directory": project.solidworks_part_directory,
            "updated_at": updated,
        }

    @staticmethod
    def _schema_description(updated: str) -> dict:
        return {
            "schema_version": __version__,
            "updated_at": updated,
            "external_access": "仅供外部软件读取；请勿直接写入 SQLite 或这些快照文件。",
            "files": {
                "项目数据快照.json": "全部项目与图号记录",
                "当前项目状态.json": "当前项目、模块及最近编号；不存在项目级单一下一图号",
                "数据格式说明.json": "开放数据格式说明",
            },
            "types": {"id": "integer", "drawing_number": "integer", "archived": "boolean", "timestamps": "ISO 8601 string"},
        }

    def _atomic_json(self, name: str, data: dict) -> None:
        target = self.directory / name
        fd, temporary = tempfile.mkstemp(prefix=f".{name}.", suffix=".tmp", dir=self.directory)
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
                json.dump(data, stream, ensure_ascii=False, indent=2)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, target)
        finally:
            remove_file(temporary, what="导出临时文件")
